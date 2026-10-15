"""Build MidtermReflection_SanchezSalazar.docx for Navigate upload."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "submissions"
    / "MidtermReflection_SanchezSalazar.docx"
)


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Montserrat"
    run.font.size = Pt(12)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Montserrat"
    run.font.size = Pt(10)


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Montserrat"
    style.font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("Midterm Reflection: Security Automation With Python")
    tr.bold = True
    tr.font.name = "Montserrat"
    tr.font.size = Pt(16)

    add_heading(doc, "Approach And Development:")
    add_body(
        doc,
        "My approach to developing the socket conection and port scanner for this "
        "midterm was focused around understanding how low level communication works, "
        "being able to do proper diagnosing without overflowing the target, specialy "
        "learning how to do it in Python. For Part 1, I built a basic client-server "
        "setup with Python's built-in socket library. I wanted a reliable TCP "
        "connection, proper encoding/decoding, and to not leave sockets hanging when "
        "something went wrong. For Part 2 I made a port scanner that loops through a "
        "port range on a target. I used settimeout so the script wouldnt freeze on "
        "closed ports, and connect_ex() so closed ports just give back an error code "
        "instead of crashing everything.",
    )

    add_heading(doc, "Challenges And AI Tool Assistance:")
    add_body(
        doc,
        "One of the harder parts was dealing with connection exceptions and making "
        "sure sockets actually close when things fail, like if the client tries to "
        "connect and the server isnt even running. AI helped with some of the "
        "boilerplate and error handling. Prompts I used were like: \"Write a Python "
        "script for a basic TCP socket client and server with error handling\" and "
        "\"How do I implement a fast timeout in a Python socket port scanner?\" I "
        "didnt just copy the output though. I tested it localy, changed the timeout "
        "values, and adjusted the print statements so it matched what I needed for "
        "my screenshots and test cases.",
    )

    add_heading(doc, "Security Considerations And Best Practices:")
    add_body(
        doc,
        "When your writing scripts that talk to sockets or scan ports, a few things "
        "matter. I used try/except/finally so sockets get closed and your not "
        "leaking resources. I also checked that the port range stays between 1 and "
        "65535 before scanning. I kept the scans limited and put small delays "
        "between attempts so I wasnt just hammering the target.",
    )

    add_heading(doc, "Legal And Ethical Considerations:")
    add_body(
        doc,
        "Port scanning can get you in trouble if you point it at random systems. "
        "Under the Computer Fraud and Abuse Act (CFAA), scanning systems without "
        "permission can lead to legal issues. I only scanned approved targets like "
        "localhost (127.0.0.1) and scanme.nmap.org, which is meant for this kind of "
        "testing. Same automation skills used for security work can be misused "
        "pretty easy if your not careful about scope.",
    )

    add_heading(doc, "Reference:")
    add_body(
        doc,
        "Nmap. (n.d.). Nmap Network Scanning: The Official Nmap Reference Guide. "
        "Insecure.Com LLC. https://nmap.org/book/man.html",
    )

    add_heading(doc, "Conclusion And Key Takeaways:")
    add_body(
        doc,
        "Building the client, server, and scanner myself helped me understand TCP "
        "and sockets a lot more then just using a premade tool. It also showed me "
        "why security automation is useful, you can take these same ideas and turn "
        "them into small tools for recon or lab testing. If I did this again I "
        "would clean up the output formatting more and maybe add a simple progress "
        "indicator for longer scans.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
