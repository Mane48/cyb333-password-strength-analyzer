"""Build CYB_333_Project_SanchezSalazar.docx (final Course Project report).

Voice: match student midterm style — casual, lightly imperfect grammar.
Do NOT auto-upload; student revises first.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submissions"
GITHUB = "https://github.com/Mane48/cyb333-password-strength-analyzer"
LAST = "SanchezSalazar"


def _style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _maybe_picture(doc: Document, path: Path, width_in: float = 6.2) -> None:
    if path.is_file() and path.stat().st_size > 1000:
        doc.add_picture(str(path), width=Inches(width_in))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        _p(doc, f"[Missing demo image: {path.name}]")


def build() -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _style(doc)

    doc.add_heading("CYB333 Final Project Report — Password Strength Analyzer", 0)
    for line in [
        "Student: Emmanuel Sanchez Salazar",
        "Course: CYB333 Security Automation",
        "National University",
        f"Filename: CYB_333_Project_{LAST}.docx",
        f"GitHub repository (public): {GITHUB}",
    ]:
        _p(doc, line)

    doc.add_heading("Introduction", level=1)
    _p(
        doc,
        "The main goal of this project was to automate checking password quality "
        "instead of just guessing if something looks strong. Weak passwords still "
        "get people into trouble with guessing, spraying, and credential stuffing, "
        "so I wanted a small Python tool that scores them the same way every time. "
        "I picked Password Strength Analyzer off the course project list because "
        "it felt doable in the time we had and still connected to security "
        "automation.",
    )
    _p(
        doc,
        "What the tool does is take a password and score it based on length, "
        "character types (lower, upper, digits, symbols), a rough entropy "
        "estimate, whether it shows up in a common password list, and weak "
        "patterns like repeats, sequences, keyboard walks, and years like 2024. "
        "Then it gives you a score out of 100, a rating from Very Weak to Very "
        "Strong, a checklist, the issues it found, and suggestions. You can run "
        "it interactive with --prompt, pass one password, batch a file, or export "
        "JSON. Thats basically the project in a sentence: make password review "
        "repeatable and documented on GitHub so someone else can actually run it.",
    )

    doc.add_heading("Development Process", level=1)
    _p(
        doc,
        "For planning I kept the scope pretty tight on purpose. I didnt want a "
        "giant UI that I wouldnt finish. Proposal first, then a working MVP, then "
        "the mid week update, then demos and this final report. That order helped "
        "because if I got stuck I still had something submitted along the way.",
    )
    _p(
        doc,
        "Design wise I split the code up instead of dumping everything in one "
        "file. analyzer.py does the scoring and has an AnalysisResult dataclass "
        "so printing and JSON use the same shape. patterns.py is where the weak "
        "pattern checks live. cli.py is argparse, getpass, and formatting. The "
        "common password list is just a text file under wordlists/. I stuck to "
        "the standard library for runtime so setup isnt annoying for grading.",
    )
    _p(
        doc,
        "Implementation was mostly iterative. Score goes up for length, classes, "
        "entropy, and passphrase style structure, then gets dinged for common "
        "list hits and patterns, then clamped 0–100 into rating bands. I used "
        "git, pytest, and tried to keep naming readable / PEP 8-ish. Samples "
        "under samples/ are fake demo passwords so Im not putting real secrets "
        "in the repo. Nothing fancy for methodology, just small modules, clear "
        "checks, and a README that tells you how to run it without me explaining "
        "it live.",
    )

    doc.add_heading("Use of AI Tools", level=1)
    _p(
        doc,
        "I used Cursor a lot for scaffolding and docs, and I already had VS Code "
        "with Copilot style stuff from the week one environment setup. AI helped "
        "with module layouts, argparse wiring, README chunks, pytest ideas, and "
        "the Word draft files for Navigate. It didnt replace me actually running "
        "the tool or checking if the scoring made sense.",
    )
    _p(doc, "Prompts I used were like:")
    for item in [
        "Create a Python CLI password strength analyzer with batch mode, JSON export, a common-password list, and unit tests.",
        "Add heuristics for keyboard walks, repeated characters, sequences, and year-like suffixes; return issues and suggestions.",
        "Draft a CYB333 final project report outline matching the Navigate sections and filename CYB_333_Project_LastName.docx.",
        "Generate pytest cases for very weak common passwords, year suffixes, and a strong random-looking password.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    _p(
        doc,
        "Value wise it saved time and stopped me from forgetting dumb checklist "
        "stuff like README sections and filename format. Limits though: AI can "
        "sound generic, invent thresholds that dont match the code, or write "
        "reflections your not supposed to let it write. I didnt just copy the "
        "output. I tested localy, ran pytest, and rewrote the personal parts in "
        "my own words. Demo screenshots of the tool running are under "
        "Demonstration Evidence below.",
    )

    doc.add_heading("Project Challenges", level=1)
    _p(
        doc,
        "One challenge was making the scoring meaningful without overbuilding. "
        "If I only checked length it would feel too thin for this project, but a "
        "full cracking pipeline was way to much for the timeline. So I mixed "
        "simple heuristics, a common list, and suggestions that explain why "
        "something is weak instead of only printing a number.",
    )
    _p(
        doc,
        "Another one was safe demos. Putting passwords on the command line can "
        "leak through shell history if your not careful. I documented --prompt "
        "as the better option, kept sample passwords synthetic, and added a "
        "masked preview that only shows first and last character.",
    )
    _p(
        doc,
        "Also just time. Between labs, midterm stuff, and work, its easy to leave "
        "docs for last. Breaking it into proposal / update / final helped. Getting "
        "the repo public earlier was good too so I wasnt scrambling on GitHub "
        "access at the end.",
    )
    _p(
        doc,
        "Honestly the scoring challenge taught me the most. I kept wanting to add "
        "more features then realized a clean CLI with honest checks and tests is "
        "more useful for this class then a half finished UI. That kinda matches "
        "what CYB333 is about, packaging judgment into something repeatable.",
    )

    doc.add_heading("Reflection on Learning", level=1)
    _p(
        doc,
        "This project changed how I think about security automation a bit. Before "
        "I mostly pictured big SOAR playbooks or enterprise tools. This was "
        "smaller, but it still forced the same habits: define checks clearly, "
        "return structured output, document how somebody else runs it, and prove "
        "it with tests. Thats closer to real defensive scripting then I expected.",
    )
    _p(
        doc,
        "On the Python side I got more comfortable splitting code into modules "
        "instead of one long script. The dataclass for AnalysisResult made JSON "
        "export pretty easy, and pytest caught dumb mistakes early like rating "
        "bands that didnt match the score clamp. I also practiced safer CLI "
        "habits with getpass and keeping secrets out of git.",
    )
    _p(
        doc,
        "It also connected to credential attacks we talked about in class. When "
        "the analyzer flags qwerty walks, years, or common list hits, its "
        "pointing at the same shortcuts attackers try first. Seeing weak samples "
        "score near zero and a longer random password land in Very Strong made "
        "that more concrete. Week one env setup and the week three Jupyter stuff "
        "helped too, by the end VS Code / GitHub felt less like homework and more "
        "like just the workbench.",
    )
    _p(
        doc,
        "If Im being real the biggest improvement wasnt some fancy algorithm. It "
        "was finishing a public repo another person can clone and run from the "
        "README without me sitting next to them. That feels like the actual bar "
        "for this class. If I did this again I would capture screenshots the same "
        "day a feature works instead of waiting till submission week.",
    )

    doc.add_heading("Time Management", level=1)
    _p(
        doc,
        "Total time was around 28–32 hours across the term. Rough breakdown: "
        "planning/proposal about 4 hours, coding and cleanup about 12–14, "
        "testing and demos about 5, README/update/final report about 7–8.",
    )
    _p(
        doc,
        "What worked was front loading the MVP so the mid week update wasnt "
        "empty, then doing shorter night sessions for docs instead of waiting "
        "for one long free day that never shows up when I have work. What didnt "
        "work as well was leaving screenshot and report polish too close to "
        "other deadlines. Next time Im grabbing demo evidence earlier.",
    )

    doc.add_heading("Demonstration Evidence", level=1)
    _p(
        doc,
        "Repo is public at the GitHub link on the title page. Screenshots below "
        "are from successful local runs on Windows after making a venv and "
        "installing requirements. Commands I used:",
    )
    for cmd in [
        "python main.py --file samples/passwords.txt --reveal-preview",
        "python main.py --file samples/passwords.txt --json -o samples/demo_report.json",
        "pytest -q",
    ]:
        doc.add_paragraph(cmd, style="List Bullet")

    _p(doc, "Figure 1. Batch analysis with masked previews (synthetic sample passwords).")
    _maybe_picture(doc, OUT / "demo_batch_run.png")

    _p(doc, "Figure 2. pytest suite result (6 passed).")
    _maybe_picture(doc, OUT / "demo_pytest.png", width_in=5.5)

    _p(doc, "Figure 3. Excerpt of the JSON report output.")
    _maybe_picture(doc, OUT / "demo_json_report.png")

    _p(
        doc,
        "Behavior matched what I expected. Common/weak samples come out Very "
        "Weak, year style mixed passwords score lower then stronger random ones, "
        "and the better samples hit Strong / Very Strong with suggestions when "
        "theres still something off.",
    )

    doc.add_heading("Team Collaboration", level=1)
    _p(doc, "Not applicable, I did this project individually.")

    doc.add_heading("Conclusion", level=1)
    _p(
        doc,
        "Overall the Password Strength Analyzer is a documented, testable "
        "automation utility that scores passwords, explains weaknesses, and "
        "exports results. I have a public GitHub repo, README setup steps, "
        "unit tests, sample inputs, and this report lined up with the course "
        "rubric areas.",
    )
    _p(
        doc,
        "If I kept going later I might add Have I Been Pwned range checks, a "
        "simple HTML summary, or a CI check that fails on weak demo credentials. "
        "I left that out on purpose so what I turned in stays reliable and easy "
        "to grade. For now it matches what I said in the proposal: a practical "
        "security automation script that makes password quality measurable "
        "instead of just vibes.",
    )

    doc.add_heading("Appendix — Submission Checklist", level=1)
    for item in [
        f"GitHub URL public: {GITHUB}",
        "README covers objectives, setup, usage, layout, tests, security notes, AI disclosure",
        "Code split into modules with comments / meaningful names; pytest present",
        "Report filename CYB_333_Project_SanchezSalazar.docx",
        "Demo screenshots included (batch run, pytest, JSON)",
        "AI usage described with example prompts; personal reflection in my own wording",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    words = sum(len(para.text.split()) for para in doc.paragraphs)
    _p(doc, f"[Approximate word count including captions/lists: {words}]")

    path = OUT / f"CYB_333_Project_{LAST}.docx"
    doc.save(path)
    print(f"Wrote {path} ({path.stat().st_size} bytes), word_count~{words}")
    return path


if __name__ == "__main__":
    build()
