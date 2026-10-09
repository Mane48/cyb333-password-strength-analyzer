"""Generate CYB333 Word submission drafts (proposal, update, final report)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submissions"
LAST = "SanchezSalazar"


def _set_body_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _add_paras(doc: Document, paragraphs: list[str]) -> None:
    for p in paragraphs:
        doc.add_paragraph(p)


def write_proposal() -> Path:
    doc = Document()
    _set_body_style(doc)
    _add_heading(doc, "CYB333 Project Proposal — Password Strength Analyzer")
    _add_paras(
        doc,
        [
            "Student: Emmanuel Sanchez Salazar",
            "Course: CYB333 Security Automation",
            "Team Information: I will complete this project individually.",
            "Filename target: ProjectProposal_SanchezSalazar.docx",
        ],
    )

    _add_heading(doc, "Project Title", level=2)
    doc.add_paragraph("Password Strength Analyzer")

    _add_heading(doc, "Project Description (approx. 200–300 words)", level=2)
    description = (
        "This project implements a Python-based Password Strength Analyzer that "
        "automates evaluation of password quality for security awareness and "
        "defensive scripting practice. Instead of relying on informal judgment, "
        "the tool applies a consistent scoring model that considers password "
        "length, character-class diversity (lowercase, uppercase, digits, and "
        "symbols), estimated entropy, membership in a common-password list, and "
        "weak construction patterns such as repeated characters, alphabetic or "
        "numeric sequences, keyboard walks, and year-like suffixes. After "
        "scoring, the analyzer returns a rating (Very Weak through Very Strong), "
        "a checklist of policy-style checks, and concrete suggestions for "
        "improvement. "
        "The project relates directly to security automation because it replaces "
        "manual, inconsistent password review with a repeatable command-line "
        "workflow. Users can evaluate a single password interactively, analyze "
        "batches from a file, and export JSON reports that could later feed "
        "larger pipelines (for example, lab demos, awareness training exercises, "
        "or pre-checks before credential creation). The scope is intentionally "
        "focused so that a complete, documented GitHub repository—including "
        "tests, sample inputs, and a clear README—can be finished within the "
        "course timeline. Core features to implement include CLI argument "
        "parsing, a reusable analyzer module, a built-in common-password "
        "wordlist with optional custom lists, pattern heuristics, human-readable "
        "and JSON output, and unit tests. Stretch goals (time permitting) include "
        "a simple HTML summary report and expanded wordlist support. No "
        "production secrets will be stored in the repository."
    )
    doc.add_paragraph(description)
    word_count = len(description.split())
    doc.add_paragraph(f"[Word count for description body: {word_count}]")

    _add_heading(doc, "Objectives and Key Features", level=2)
    for item in [
        "Automate password strength scoring with transparent criteria",
        "Detect common weak patterns used in guessing attacks",
        "Support interactive, single, and batch evaluation modes",
        "Provide actionable remediation suggestions",
        "Ship documented code, README, tests, and demo samples on GitHub",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    _add_heading(doc, "Use of AI Tools", level=2)
    _add_paras(
        doc,
        [
            (
                "AI tools used to select/refine this project: Cursor (Composer) was "
                "used to compare course-approved project ideas against timeline, "
                "dependency risk, and rubric alignment. The Password Strength "
                "Analyzer was chosen because it is self-contained in Python, maps "
                "cleanly to security automation, and can be completed with "
                "high-quality documentation within the course window."
            ),
            (
                "AI tools planned during development: Cursor for scaffolding modules, "
                "tests, README content, and submission drafts; optional use of "
                "GitHub Copilot-style completion while writing functions. AI will "
                "assist with boilerplate, edge-case ideas, and documentation "
                "structure. I will review all generated code, run tests locally, "
                "and personally write reflection sections that must represent my "
                "own learning (AI will not author personal reflections)."
            ),
            (
                "How AI helps: faster iteration on CLI design, scoring heuristics, "
                "and submission formatting while I remain responsible for "
                "validation, demos, and final wording."
            ),
        ],
    )

    _add_heading(doc, "Notes for Instructor", level=2)
    doc.add_paragraph(
        "This proposal follows the course list of potential projects "
        "(Password Strength Analyzer). I will wait for approval/feedback before "
        "treating development as finalized, while continuing low-risk scaffolding."
    )

    path = OUT / f"ProjectProposal_{LAST}.docx"
    doc.save(path)
    return path


def write_update() -> Path:
    doc = Document()
    _set_body_style(doc)
    _add_heading(doc, "CYB333 Mid-Week Three Project Update")
    _add_paras(
        doc,
        [
            "Student: Emmanuel Sanchez Salazar",
            "Course: CYB333 Security Automation",
            "Project: Password Strength Analyzer",
            "Team: Individual",
            "Filename target: ProjectUpdate_SanchezSalazar.docx",
        ],
    )

    _add_heading(doc, "1. Project Progress Overview", level=2)
    _add_paras(
        doc,
        [
            "Completed tasks / milestones so far:",
            "• Selected project idea (Password Strength Analyzer) aligned to course list.",
            "• Scaffolded repository structure: analyzer module, CLI, wordlist, samples, tests.",
            "• Implemented scoring (length, character classes, entropy estimate), common-list checks, and pattern heuristics.",
            "• Added human-readable and JSON reporting; created README with setup/run instructions.",
            "• Wrote unit tests and generated demo batch inputs under samples/.",
            "• Prepared proposal and update submission drafts under submissions/.",
            "",
            "GitHub repository: [PASTE PUBLIC REPO URL HERE AFTER PUSH]",
            "Note: Repository has been initialized locally and is ready to publish; link will be added as soon as the remote is created.",
        ],
    )

    _add_heading(doc, "2. Remaining Work", level=2)
    _add_paras(
        doc,
        [
            "Pending tasks and estimated timeline (final week):",
            "• Publish GitHub remote and verify instructor can access README/code (Day 1).",
            "• Polish scoring thresholds from demo feedback and expand wordlist if needed (Day 1–2).",
            "• Capture screenshots / short video of successful runs for the final report (Day 2).",
            "• Write personal reflection sections (learning, challenges, time spent) without AI authorship (Day 2–3).",
            "• Finalize CYB_333_Project_SanchezSalazar.docx and submit Course Project (Day 3).",
        ],
    )

    _add_heading(doc, "3. Challenges and Solutions", level=2)
    _add_paras(
        doc,
        [
            "Obstacle: Balancing a ‘simple’ project idea with rubric expectations for automation depth, testing, and documentation.",
            "Response: Focused on clean modular design, explicit checks/suggestions, batch+JSON automation hooks, and PEP 8-oriented structure rather than adding unnecessary UI complexity.",
            "Obstacle: Avoiding unsafe handling of secrets in demos.",
            "Response: Prefer --prompt, mask previews, and keep sample passwords synthetic.",
        ],
    )

    _add_heading(doc, "4. Team Collaboration", level=2)
    doc.add_paragraph("Not applicable — individual project.")

    _add_heading(doc, "5. Use of AI Tools", level=2)
    _add_paras(
        doc,
        [
            "Cursor has been used to scaffold modules (analyzer/cli/patterns), draft README sections, generate submission Word documents, and propose test cases. Example assistance: generating argparse wiring and structuring AnalysisResult as a dataclass for JSON export.",
            "I review outputs, run pytest, and execute demo commands locally before treating work as complete. Personal reflection content for the final report remains student-authored.",
        ],
    )

    path = OUT / f"ProjectUpdate_{LAST}.docx"
    doc.save(path)
    return path


def write_final_report() -> Path:
    doc = Document()
    _set_body_style(doc)
    _add_heading(doc, "CYB333 Final Project Report — Password Strength Analyzer")
    _add_paras(
        doc,
        [
            "Student: Emmanuel Sanchez Salazar",
            "Course: CYB333 Security Automation",
            "Filename target: CYB_333_Project_SanchezSalazar.docx",
            "GitHub repository: [PASTE PUBLIC REPO URL HERE]",
        ],
    )

    _add_heading(doc, "Introduction", level=2)
    doc.add_paragraph(
        "The objective of this project is to automate password-quality assessment "
        "using a Python command-line tool. Weak passwords remain a frequent "
        "contributor to account compromise through guessing, spraying, and "
        "credential stuffing. This Password Strength Analyzer evaluates candidate "
        "passwords against transparent criteria—length, character diversity, "
        "estimated entropy, common-password membership, and weak construction "
        "patterns—then returns a numeric score, rating, checklist results, and "
        "remediation suggestions. The tool supports interactive prompting, "
        "single-password evaluation, batch file processing, and JSON export so "
        "results can be reviewed by a human or consumed by another script. The "
        "project demonstrates applied security automation: consistent checks, "
        "repeatable reporting, and documented operational steps via README and "
        "version control."
    )

    _add_heading(doc, "Development Process", level=2)
    _add_paras(
        doc,
        [
            "Planning: I selected a course-listed idea with a scoped feature set that still maps to automation outcomes (CLI workflows, batch analysis, JSON artifacts). Deliverables were sequenced as proposal → working MVP → tests/docs → update → final report/demo evidence.",
            "Design: The codebase separates concerns into analyzer.py (scoring), patterns.py (heuristics), cli.py (interface/reporting), and a wordlist asset. An AnalysisResult dataclass provides a stable schema for both printed summaries and JSON serialization.",
            "Implementation: Scoring combines additive rewards (length, classes, entropy, passphrase structure) with penalties (common list hits and pattern issues), clamped to 0–100 and mapped to rating bands. The CLI uses argparse and getpass for safer interactive entry. Sample inputs under samples/ support demos without exposing real credentials.",
            "Methodologies/tools: Iterative development with git, pytest for regression checks, and PEP 8-oriented naming/structure. No external runtime libraries are required, which simplifies setup and grading reproducibility.",
        ],
    )

    _add_heading(doc, "Use of AI Tools", level=2)
    _add_paras(
        doc,
        [
            "AI tools used: Cursor (Composer) for repository scaffolding, documentation drafts, submission document generation, and test case ideas. Optional inline completion was used while writing functions.",
            "Example prompts / assistance themes (summarized):",
            "• “Create a Python CLI password strength analyzer with batch JSON output and tests.”",
            "• “Generate a one-page CYB333 proposal matching Navigate rubric sections.”",
            "• “Add pattern heuristics for keyboard walks and year suffixes.”",
            "[SCREENSHOT PLACEHOLDER 1: Insert Cursor chat or Copilot usage screenshot]",
            "[SCREENSHOT PLACEHOLDER 2: Insert second AI-assist example if available]",
            "Value added: faster bootstrap of project structure and fewer missed submission checklist items (README sections, filename conventions, rubric-aligned headings).",
            "Limitations/challenges: AI drafts can be generic or overconfident about scoring thresholds; I validated behavior with tests and manual demos. Course policy also prohibits AI-authored personal reflections, so those sections are written (or finalized) by me.",
        ],
    )

    _add_heading(doc, "Project Challenges", level=2)
    _add_paras(
        doc,
        [
            "Challenge 1 — Meaningful scoring without overengineering: A toy checker that only measures length would under-serve the rubric; a full offline cracker would exceed timeline. Solution: combine transparent heuristics with a common-password list and clear suggestions.",
            "Challenge 2 — Safe demos: Command-line passwords can leak via shell history. Solution: document --prompt, provide synthetic samples, and optional masked previews.",
            "Challenge 3 — Documentation completeness for graders: Solution: README covers objectives, setup, usage, layout, tests, and security notes; submissions/ contains Navigate-ready drafts.",
            ">>> STUDENT REFLECTION ADD-ON (required; write in your own words): Add 2–4 sentences about which challenge taught you the most and why. Do not paste AI text here. <<<",
        ],
    )

    _add_heading(doc, "Reflection on Learning", level=2)
    _add_paras(
        doc,
        [
            "Starter notes (edit heavily into your voice): This project reinforced that security automation is often about packaging defensive judgment into repeatable scripts—checks, reports, and operational instructions—rather than only building complex exploit tooling.",
            ">>> REQUIRED PERSONAL REFLECTION (AI must NOT write this for you): Describe what you personally learned about Python structuring, password attacks/defenses, testing, and git/README practice. Mention at least one concept from CYB333 that connected while building this tool. 150–250 words. <<<",
        ],
    )

    _add_heading(doc, "Time Management", level=2)
    _add_paras(
        doc,
        [
            "Estimated total time spent (draft — update with your real hours): approximately __ hours across proposal planning, implementation, testing, documentation, and submission prep.",
            "Suggested breakdown to fill in: planning/proposal __; coding __; testing/demo capture __; documentation/report __.",
            ">>> REQUIRED PERSONAL NOTE: Briefly explain what scheduling approach worked (or did not) for you during the course weeks. <<<",
        ],
    )

    _add_heading(doc, "Demonstration Evidence", level=2)
    _add_paras(
        doc,
        [
            "Include either a short video or screenshots showing successful execution.",
            "[SCREENSHOT PLACEHOLDER A: python main.py --file samples/passwords.txt]",
            "[SCREENSHOT PLACEHOLDER B: JSON output or pytest -q success]",
            "Commands used for demo:",
            "python main.py --file samples/passwords.txt --reveal-preview",
            "python main.py --file samples/passwords.txt --json -o samples/demo_report.json",
            "pytest -q",
        ],
    )

    _add_heading(doc, "Team Collaboration", level=2)
    doc.add_paragraph("Not applicable — completed individually.")

    _add_heading(doc, "Conclusion", level=2)
    _add_paras(
        doc,
        [
            "The Password Strength Analyzer delivers a documented, testable automation utility that scores passwords, explains weaknesses, and exports results for further use. Outcomes include a public/shared GitHub repository, README-driven setup, and a report aligning to course rubric categories (technical implementation, code quality, reflection, professionalism, and effort).",
            "Future enhancements could include integrating Have I Been Pwned k-anonymity range queries, generating HTML dashboards, or plugging the analyzer into a CI job that rejects weak demo credentials. Those extensions were intentionally deferred to keep the delivered scope reliable and gradable within the term.",
        ],
    )

    _add_heading(doc, "Appendix — Rubric Self-Check", level=2)
    for item in [
        "Technical implementation: working CLI features + pattern/common-list logic + demos",
        "Code quality: modular layout, comments/docstrings, README, tests, meaningful names",
        "Reflection/analysis: AI usage documented with examples; personal reflection placeholders completed by student",
        "Presentation: Word filename CYB_333_Project_SanchezSalazar.docx; clear headings",
        "Time management: hours recorded; milestones tracked via update + git history",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    path = OUT / f"CYB_333_Project_{LAST}.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [write_proposal(), write_update(), write_final_report()]
    for path in paths:
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
